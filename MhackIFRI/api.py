from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from typing import List
import os
import io
import re
import json
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.chains import create_retrieval_chain
from langchain.chains.combine_documents import create_stuff_documents_chain
from langchain.prompts import ChatPromptTemplate

# Importer OpenAIEmbeddings depuis le package langchain-openai pour éviter les avertissements
from langchain_openai import OpenAIEmbeddings
from langchain.vectorstores import Chroma
from langchain.llms import OpenAI

from langchain_google_genai import ChatGoogleGenerativeAI

from langchain_chroma import Chroma
from langchain_google_genai import GoogleGenerativeAIEmbeddings

from dotenv import load_dotenv
load_dotenv()

# Pour la gestion de la clé API
from gemin import keygemini
os.environ["GOOGLE_API_KEY"] = keygemini

app = FastAPI()

# Fonction pour identifier le type de fichier
def identify_file_type(file: UploadFile):
    ext = os.path.splitext(file.filename)[-1].lower()
    return ext

# Extraction de texte d'un fichier PDF
def extract_text_from_pdf(file: UploadFile) -> str:
    reader = PdfReader(file.file)
    extracted_text = ""
    for page in reader.pages:
        extracted_text += page.extract_text() or ""
    return extracted_text

# Extraction de texte d'un fichier DOCX
def extract_text_from_docx(file: UploadFile) -> str:
    doc = Document(file.file)
    extracted_text = ""
    for para in doc.paragraphs:
        extracted_text += para.text + "\n"
    return extracted_text

# Extraction de texte d'un fichier PPTX
def extract_text_from_pptx(file: UploadFile) -> str:
    presentation = Presentation(file.file)
    extracted_text = ""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                extracted_text += shape.text + "\n"
    return extracted_text

# Initialisation de LangChain (global pour être utilisé par /query/)
text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000)
embeddings = GoogleGenerativeAIEmbeddings(model="models/embedding-001")
llm =  ChatGoogleGenerativeAI(model="gemini-1.5-pro",temperature=0.3, max_tokens=1024000)

system_prompt = (
    "Vous êtes un assistant spécialisé dans la génération de QCM à partir du contexte fourni. "
    "Votre tâche est de produire une liste de questions sous forme de JSON structuré.\n\n"
    "Le format attendu est :\n"
    "{{\n"
    "  \"questions\": [\n"
    "    {{\n"
    "      \"id\": <numéro de la question>,\n"
    "      \"question\": \"<texte de la question>\",\n"
    "      \"choices\": {{\n"
    "        \"A\": \"<option A>\",\n"
    "        \"B\": \"<option B>\",\n"
    "        \"C\": \"<option C>\",\n"
    "        \"D\": \"<option D>\"\n"
    "      }},\n"
    "      \"answer\": \"<bonne réponse (ex: A, B, C ou D)>\",\n"
    "      \"explanation\": {{\n"
    "        \"A\": \"<explication pour A>\",\n"
    "        \"B\": \"<explication pour B>\",\n"
    "        \"C\": \"<explication pour C>\",\n"
    "        \"D\": \"<explication pour D>\"\n"
    "      }}\n"
    "    }}\n"
    "  ]\n"
    "}}\n\n"
    "Générez plusieurs QCM en respectant ce format. "
    "Si le contexte ne permet pas de générer un QCM valide, retournez un JSON vide : {{ \"questions\": [] }}.\n"
    "Le contexte est le suivant :\n{context}"
)


prompt = ChatPromptTemplate.from_messages(
    [
        ("system", system_prompt),
        ("human", "{input}"),
    ]
)


# Créer la chaîne de question-réponse
question_answer_chain = create_stuff_documents_chain(llm, prompt)
rag_chain = None  # Ce global sera initialisé après upload

# Point de terminaison pour télécharger et traiter un fichier
@app.post("/upload/")
async def upload_file(file: UploadFile = File(...)):
    ext = identify_file_type(file)
    # Réinitialiser le curseur du fichier pour éviter tout problème de lecture
    file.file.seek(0)
    if ext == '.pdf':
        extracted_text = extract_text_from_pdf(file)
    elif ext == '.docx':
        extracted_text = extract_text_from_docx(file)
    elif ext == '.pptx':
        extracted_text = extract_text_from_pptx(file)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type")

    # Traitement du texte extrait avec LangChain
    docs = text_splitter.split_text(extracted_text)
    vectorstore = Chroma.from_texts(docs, embeddings)
    retriever = vectorstore.as_retriever(search_type="similarity", search_kwargs={"k": 10})

    global rag_chain
    rag_chain = create_retrieval_chain(retriever, question_answer_chain)

    return JSONResponse(content={"text": extracted_text})

# Modèle de requête pour interroger le LLM

# Fonction pour nettoyer la chaîne JSON en retirant les caractères de contrôle invalides
def sanitize_json_string(s: str) -> str:
    # Retirer les balises de code markdown s'il y en a
    s = s.strip()
    if s.startswith("```"):
        s = "\n".join(s.split("\n")[1:])
        if s.endswith("```"):
            s = "\n".join(s.split("\n")[:-1])
    # Supprimer les caractères de contrôle non désirés (sauf tabulation, \n et \r)
    s = re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]', '', s)
    # Supprimer les virgules finales avant } ou ]
    s = re.sub(r",\s*([\]}])", r"\1", s)
    return s

class QueryRequest(BaseModel):
    query: str

@app.post("/query/")
async def query_langchain(request: QueryRequest):
    if rag_chain is None:
        raise HTTPException(status_code=400, detail="No document has been uploaded yet.")
    try:
        response = rag_chain.invoke({"input": request.query})
    except Exception as e:
        return JSONResponse(status_code=500, content={
            "error": "Erreur lors de l'exécution de rag_chain.invoke",
            "detail": str(e)
        })

    raw_answer = response.get("answer", "")
    if not raw_answer:
        return JSONResponse(status_code=500, content={
            "error": "Réponse vide de rag_chain.invoke"
        })
    
    # Nettoyer la réponse brute
    sanitized_answer = sanitize_json_string(raw_answer)
    
    try:
        parsed_json = json.loads(sanitized_answer)
    except json.JSONDecodeError as e:
        # Si l'erreur persiste, on renvoie l'erreur et la réponse brute pour débogage
        return JSONResponse(status_code=500, content={
            "error": "Erreur JSON après nettoyage",
            "detail": str(e),
            "raw_response": raw_answer
        })

    return {"answer": parsed_json}

 